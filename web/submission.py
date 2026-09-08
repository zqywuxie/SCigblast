"""Immutable XLSX revisions; Excel cell coordinates are retained for editing."""
from __future__ import annotations

import json
import re
import shutil
import uuid
import zipfile
from pathlib import Path

from openpyxl import load_workbook

MAX_BYTES = 20 * 1024 * 1024


def checked_workbook(path):
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("XLSX exceeds 20 MB")
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Expanded XLSX exceeds 100 MB")
    book = load_workbook(path)
    if sum(s.max_row * s.max_column for s in book) > 500_000:
        book.close()
        raise ValueError("Workbook exceeds 500,000 cells")
    return book


def resolve(root: Path, token: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}/[a-f0-9]{32}", token):
        raise ValueError("Invalid submission revision")
    folder = root / "submissions" / token
    if not (folder / "metadata.json").is_file():
        raise ValueError("Submission revision not found")
    return folder


def field_type(value):
    text = str(value or "").lower()
    if "note" in text or "储存地址" in text:
        return "note"
    if "sample id" in text:
        return "sample_id"
    if "dual index" in text:
        return "dual_index"
    if "barcode" in text:
        return "barcode"
    if "chain" in text:
        return "chain"
    if "species" in text or "speicies" in text:
        return "species"
    return ""


def inspect(root: Path, token: str):
    folder = resolve(root, token)
    metadata = json.loads((folder / "metadata.json").read_text(encoding="utf-8"))
    sheets = []
    for file in sorted(folder.glob("*.xlsx")):
        book = checked_workbook(file)
        try:
            for sheet in book:
                header_row = next((r for r in range(1, min(sheet.max_row, 30) + 1)
                                   if any(field_type(c.value) == "sample_id" for c in sheet[r])), None)
                if header_row is None:
                    continue
                headers = [str(c.value or "") for c in sheet[header_row]]
                types = [field_type(h) for h in headers]
                rows, current_note, note_anchor = [], "", None
                for cells in sheet.iter_rows(min_row=header_row + 1):
                    if not any(c.value is not None for c in cells):
                        continue
                    values, editable = [], []
                    for index, c in enumerate(cells):
                        kind = types[index]
                        value = "" if c.value is None else str(c.value)
                        anchor = c.coordinate
                        merge = next((m for m in sheet.merged_cells.ranges if c.coordinate in m), None)
                        if merge:
                            anchor = sheet.cell(merge.min_row, merge.min_col).coordinate
                            value = str(sheet[anchor].value or "")
                        if kind == "note":
                            if value:
                                current_note, note_anchor = value, anchor
                            elif note_anchor:
                                value, anchor = current_note, note_anchor
                        values.append(value)
                        editable.append({"cell": anchor, "kind": kind,
                                         "merged": str(merge) if merge else "",
                                         "inherited": anchor != c.coordinate})
                    rows.append({"row": cells[0].row, "values": values, "editable": editable})
                sheets.append({"file": file.name, "sheet": sheet.title, "headers": headers, "rows": rows})
        finally:
            book.close()
    return {"revision": token, "metadata": metadata, "sheets": sheets}


def create(root: Path, files: list[Path], source: str):
    token = f"{uuid.uuid4().hex}/{uuid.uuid4().hex}"
    folder = root / "submissions" / token
    folder.mkdir(parents=True)
    for index, path in enumerate(files):
        book = checked_workbook(path)
        book.close()
        shutil.copyfile(path, folder / f"{index + 1:03d}_{path.name}")
    (folder / "metadata.json").write_text(json.dumps({"source": source, "parent": None}, ensure_ascii=False), encoding="utf-8")
    result = inspect(root, token)
    if not result["sheets"]:
        raise ValueError("No worksheet with Sample ID header found")
    return result


def revise(root: Path, token: str, changes: list[dict]):
    old = resolve(root, token)
    view = inspect(root, token)
    allowed = {(s["file"], s["sheet"], c["cell"]): c["kind"]
               for s in view["sheets"] for row in s["rows"] for c in row["editable"] if c["kind"]}
    edits = {}
    for change in changes:
        key = (change["file"], change["sheet"], change["cell"])
        if key not in allowed:
            raise ValueError("Cell is not editable")
        value = str(change["value"])
        if len(value) > 4096 or value.startswith("="):
            raise ValueError("Invalid metadata value")
        if key in edits and edits[key] != value:
            raise ValueError("Conflicting changes to the same merged cell")
        edits[key] = value
    new = f"{token.split('/')[0]}/{uuid.uuid4().hex}"
    folder = root / "submissions" / new
    folder.mkdir(parents=True)
    for file in old.glob("*.xlsx"):
        book = checked_workbook(file)
        try:
            for (filename, sheet, cell), value in edits.items():
                if filename == file.name:
                    book[sheet][cell] = value
            book.save(folder / file.name)
        finally:
            book.close()
    (folder / "metadata.json").write_text(json.dumps({"source": view["metadata"]["source"], "parent": token,
                                                    "changes": changes}, ensure_ascii=False), encoding="utf-8")
    return inspect(root, new)
