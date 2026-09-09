"""Web contract tests; no real FASTQ or pipeline commands are executed."""
import csv
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as web
import submission
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.raw = self.root / 'raw'; self.raw.mkdir()
        self.out = self.root / 'output'
        self.source = self.root / 'metadata.xlsx'
        book = Workbook(); sheet = book.active
        sheet.append(['Sample ID File new name', 'Dual Index', 'Chain', 'Note', 'Barcode'])
        sheet.append(['A', 'A01', 'T', str(self.raw), 'BC_1'])
        sheet.append(['B', 'A02', 'T', None, 'BC_2'])
        sheet.merge_cells('D2:D3'); book.save(self.source); book.close()
        self.patches = [patch.object(web, 'STATE_ROOT', self.root / 'state'),
                        patch.object(web, 'DB_PATH', self.root / 'state/db.sqlite'),
                        patch.object(web, 'PIPELINE_ROOT', Path(__file__).resolve().parents[2]),
                        patch.object(web, 'DEFAULT_OUTPUT_ROOT', self.out),
                        patch.dict(os.environ, {k: str(self.root) for k in ('SCIGBLAST_ALLOWED_INPUT_ROOTS', 'SCIGBLAST_ALLOWED_SUBMISSION_ROOTS', 'SCIGBLAST_ALLOWED_OUTPUT_ROOTS')}),
                        patch.object(web, 'schedule')]
        for p in self.patches: p.start()
        web.PROCESSES.clear(); web.init_db(); self.client = TestClient(web.app)
        self.view = submission.create(web.STATE_ROOT, [self.source], str(self.source))

    def tearDown(self):
        self.client.close(); web.PROCESSES.clear()
        for p in reversed(self.patches): p.stop()
        self.temp.cleanup()

    def create(self, dataset='batch', pipeline='igblast_base'):
        r = self.client.post('/api/jobs', json={'pipeline':pipeline, 'operator':'test', 'input_path':str(self.raw), 'submission_revision':self.view['revision'], 'output_root':str(self.out), 'dataset_label':dataset})
        self.assertEqual(r.status_code,200,r.text);return r.json()['id']

    def match(self, jid, rows=151):
        row=web.get_job_row(jid);folder=self.out/'01.match'/row['dataset'];folder.mkdir(parents=True,exist_ok=True)
        path=folder/'sample_manifest.csv'
        with path.open('w',newline='',encoding='utf-8') as h:
            w=csv.DictWriter(h,fieldnames=['sample_id','status','error','r1_path','igblast_chains']);w.writeheader()
            for i in range(rows):w.writerow({'sample_id':str(i),'status':'ERROR' if i==150 else 'OK','error':'missing barcode' if i==150 else '', 'r1_path':f'/raw/{i}_R1.fq','igblast_chains':'TRA,TRB'})
        web.update_job(jid,status='WAITING_REVIEW',attempt_no=1)
        return path

    def test_merged_note_revision_preserves_original(self):
        rows=self.view['sheets'][0]['rows'];self.assertEqual(rows[0]['values'][3],rows[1]['values'][3]);self.assertEqual(rows[1]['editable'][3]['cell'],'D2')
        dest=self.root/'new';dest.mkdir()
        change={'file':self.view['sheets'][0]['file'],'sheet':'Sheet','cell':'D2','value':str(dest)}
        r=self.client.post('/api/submissions/revise',json={'revision':self.view['revision'],'changes':[change]});self.assertEqual(r.status_code,200,r.text)
        self.assertEqual([row['values'][3] for row in r.json()['sheets'][0]['rows']],[str(dest)]*2)
        b=load_workbook(self.source);self.assertEqual(b.active['D2'].value,str(self.raw));b.close()
        download=self.client.get('/api/submissions/download',params={'revision':r.json()['revision'],'filename':change['file'],'original':'true'})
        b=load_workbook(io.BytesIO(download.content));self.assertEqual(b.active['D2'].value,str(self.raw));b.close()

    def test_image_runtime_on_resume_and_preflight(self):
        jid = self.create()
        row = web.get_job_row(jid)
        saved = json.loads(row['env_json'])
        saved['PATH'] = '/old/host/conda/bin'
        saved['SCIGBLAST_RUNTIME_BIN_DIR'] = '/old/runtime/bin'
        web.update_job(jid, status='QUEUED', env_json=json.dumps(saved))
        runtime = str(Path('/opt/conda/bin'))
        with patch.dict(os.environ, {'SCIGBLAST_RUNTIME_BIN_DIR': runtime, 'PATH': '/opt/conda/bin:/usr/bin'}), \
                patch.object(web.subprocess, 'Popen') as spawn, patch.object(web.threading, 'Thread'):
            spawn.return_value.pid = 12345
            self.assertTrue(web.launch_job(jid))
            actual = spawn.call_args.kwargs['env']
            self.assertEqual(actual['PATH'], '/opt/conda/bin:/usr/bin')
            self.assertEqual(actual['SCIGBLAST_RUNTIME_BIN_DIR'], runtime)
            with patch.object(web.shutil, 'which', side_effect=lambda command: command), \
                    patch.object(web.subprocess, 'run') as run:
                run.return_value.returncode = 0
                run.return_value.stdout = '[]'
                result = self.client.get('/api/preflight?pipeline=pig_igblast').json()
                self.assertEqual(result['commands']['igblastn'], str(Path(runtime) / 'igblastn'))

    def test_upload_and_bad_files(self):
        r=self.client.post('/api/submissions/upload?filename=test.xlsx',content=self.source.read_bytes());self.assertEqual(r.status_code,200)
        self.assertEqual(self.client.post('/api/submissions/upload?filename=test.xlsx',content=b'bad').status_code,400)
        self.assertEqual(self.client.get('/api/submissions',params={'revision':'../../etc'}).status_code,400)
        self.assertEqual(self.client.get('/api/browse',params={'path':str(self.root.parent)}).status_code,400)

    def test_complete_match_filter_review_and_stale_confirmation(self):
        jid=self.create();path=self.match(jid)
        r=self.client.get(f'/api/jobs/{jid}/match-preview',params={'errors_only':'true'}).json()
        self.assertEqual(r['total'],1);self.assertEqual(r['rows'][0]['sample_id'],'150');self.assertEqual(r['counts']['total'],151)
        body={'revision':r['revision'],'row_key':'150','label':'待补资料'}
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/review',json=body).status_code,200)
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/resume',json={}).status_code,409)
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/confirm-match',json={'revision':'old'}).status_code,409)
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/confirm-match',json={'revision':r['revision']}).status_code,200)
        self.assertEqual(json.loads(web.get_job_row(jid)['env_json'])['SCIGBLAST_WEB_MATCH_ONLY'],'0')
        self.assertEqual(self.client.post('/api/jobs',json={'pipeline':'igblast_base','operator':'test','input_path':str(self.raw),'submission_revision':self.view['revision'],'output_root':str(self.out),'dataset_label':'batch'}).status_code,409)

    def test_rematch_allows_new_but_blocks_changed_assignments(self):
        jid=self.create();p=self.match(jid,1)
        def confirm():
            web.update_job(jid,status='WAITING_REVIEW')
            revision=self.client.get(f'/api/jobs/{jid}/match-preview').json()['revision']
            return self.client.post(f'/api/jobs/{jid}/confirm-match',json={'revision':revision})
        self.assertEqual(confirm().status_code,200)
        with p.open('a') as h:h.write('new,OK,,/raw/new_R1.fq,TRA\n')
        self.assertEqual(confirm().status_code,200)
        p.write_text('sample_id,status,r1_path,igblast_chains\nchanged,OK,/raw/0_R1.fq,TRA\n',encoding='utf-8')
        self.assertEqual(confirm().status_code,409)

    def test_result_path_pig_markers_progress_and_log_rotation(self):
        jid=self.create(pipeline='pig_igblast');row=web.get_job_row(jid)
        folder=self.out/'05.igblastn_out'/'batch';folder.mkdir(parents=True)
        (folder/'chain_summary.csv').write_text('sample,chain,input_sequences,mapped_seqs,mapping_percent,filtered_seqs\nA,TRA,100,80,80,40\n',encoding='utf-8')
        result=self.client.get(f'/api/jobs/{jid}/results').json();self.assertEqual(result['rows'][0]['mapped_seqs'],'80')
        state=web.state_dir(row);state.mkdir(parents=True);(state/'.pipeline_stage_igblast.DONE').write_text('status=DONE\n')
        self.assertTrue(web.pipeline_done(row));self.assertIn('05.igblast',web.snapshot(row)['completed_stages'])
        log=web.log_path(row);log.parent.mkdir(parents=True);log.write_text('[PIG] fastp (sample-level resume enabled)\nprogress=2/10 percent=20\n[PIG] clean (sample-level resume enabled)\n')
        self.assertEqual(web.parse_progress(row),('03.clean',0))
        old=self.client.get(f'/api/jobs/{jid}/log').json();web.update_job(jid,attempt_no=2)
        new=self.client.get(f'/api/jobs/{jid}/log',params={'source':old['source'],'offset':old['next_offset']}).json();self.assertTrue(new['reset'])

    def test_scheduler_atomic_limit(self):
        ids=[self.create(f'b{i}') for i in range(4)]
        # Restore the real scheduler while replacing only process launch.
        self.patches[-1].stop()
        calls=[]
        def launch(jid):
            calls.append(jid);web.PROCESSES[jid]=object();web.update_job(jid,status='RUNNING');return True
        with patch.object(web,'launch_job',side_effect=launch):
            threads=[threading.Thread(target=web.schedule) for _ in range(8)]
            for t in threads:t.start()
            for t in threads:t.join()
        self.assertEqual(len(calls),2);self.assertEqual(len(set(calls)),2)
        self.patches[-1].start()

    def test_all_error_cannot_continue_and_queue_can_stop(self):
        jid=self.create();path=self.match(jid,1)
        path.write_text('sample_id,status,error\nA,ERROR,no match\n',encoding='utf-8')
        revision=self.client.get(f'/api/jobs/{jid}/match-preview').json()['revision']
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/confirm-match',json={'revision':revision}).status_code,409)
        web.update_job(jid,status='QUEUED')
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/stop').json()['status'],'STOPPED')

    def test_launch_failure_clears_stale_review(self):
        jid=self.create();state=web.state_dir(web.get_job_row(jid));state.mkdir(parents=True)
        marker=state/'.match_review.done';marker.write_text('status=READY_FOR_REVIEW\n')
        with patch.object(web.subprocess,'Popen',side_effect=OSError('missing bash')):
            self.assertFalse(web.launch_job(jid))
        self.assertFalse(marker.exists());self.assertEqual(web.get_job_row(jid)['status'],'FAILED');self.assertEqual(web.active_count(),0)

    def test_metadata_review_and_mode_stages(self):
        jid=self.create();self.match(jid,1)
        data=self.client.get(f'/api/jobs/{jid}/metadata-review').json()
        self.assertEqual([r['sample_id'] for r in data['rows']],['A','B'])
        row=dict(web.get_job_row(jid));row['pipeline']='ir_split';row['options_json']=json.dumps({'ir_variant':'merged'})
        self.assertEqual(web.job_config(row)['stages'][-1],'06.igblast')
        self.assertNotIn('06.representative',web.job_config(row)['stages'])

    def test_twenty_seven_row_note_group(self):
        source=self.root/'27.xlsx';b=Workbook();s=b.active
        s.append(['Sample ID File new name','Note'])
        for i in range(27):s.append([f'S{i}',str(self.raw) if i==0 else None])
        s.merge_cells('B2:B28');b.save(source);b.close()
        view=submission.create(web.STATE_ROOT,[source],str(source))
        revised=submission.revise(web.STATE_ROOT,view['revision'],[{'file':view['sheets'][0]['file'],'sheet':'Sheet','cell':'B2','value':'/analysis/new'}])
        self.assertEqual(len(revised['sheets'][0]['rows']),27)
        self.assertTrue(all(r['values'][1]=='/analysis/new' for r in revised['sheets'][0]['rows']))

    def test_batch_review_atomic_and_version_bound(self):
        jid=self.create();self.match(jid)
        before=self.client.get(f'/api/jobs/{jid}/match-preview').json()
        revision=before['revision']; url=f'/api/jobs/{jid}/review'
        self.client.post(url,json={'revision':revision,'row_key':'0','label':'待补资料','note':'keep note'})
        result=self.client.post(url,json={'revision':revision,'row_keys':['0','50','150','50'],'label':'已核对'})
        self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.json()['updated'],3)
        after=self.client.get(f'/api/jobs/{jid}/match-preview').json()
        self.assertEqual(after['revision'],revision);self.assertEqual(after['counts'],before['counts'])
        self.assertEqual(after['reviews']['0']['note'],'keep note')
        self.assertTrue(all(after['reviews'][k]['label']=='已核对' for k in ['0','50','150']))
        self.assertEqual(self.client.post(url,json={'revision':revision,'row_keys':['0','9999'],'label':'待补资料'}).status_code,400)
        self.assertEqual(self.client.get(f'/api/jobs/{jid}/match-preview').json()['reviews']['0']['label'],'已核对')
        self.assertEqual(self.client.post(url,json={'revision':'old','row_keys':['0'],'label':''}).status_code,409)
        result=self.client.post(url,json={'revision':revision,'row_keys':['0','50'],'label':''})
        self.assertEqual(result.status_code,200)
        web.update_job(jid,status='RUNNING')
        self.assertEqual(self.client.post(url,json={'revision':revision,'row_keys':['150'],'label':''}).status_code,409)

    def test_unique_default_output_names(self):
        body = {'pipeline':'igblast_base', 'operator':'郑钦云', 'input_path':str(self.raw), 'submission_revision':self.view['revision']}
        first = self.client.post('/api/jobs', json=body)
        self.assertEqual(first.status_code, 200, first.text)
        path = Path(first.json()['output_root'])
        self.assertEqual(path.parent, self.out)
        self.assertRegex(path.name, r'^郑钦云_igblast_base_\d{8}_\d{6}$')
        with patch.object(web, 'default_job_output', return_value=path):
            second = self.client.post('/api/jobs', json=body)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(second.json()['output_root'], str(path.with_name(path.name+'_02')))
        self.assertFalse(path.exists(), 'validation/creation must not create analysis files')
        with patch.object(web, 'default_job_output', return_value=self.out/'whitespace'):
            body['output_root'] = '  '
            self.assertEqual(self.client.post('/api/jobs', json=body).json()['output_root'], str(self.out/'whitespace'))

    def test_delete_exclusive_results_and_record(self):
        body = {'pipeline':'igblast_base', 'operator':'test', 'input_path':str(self.raw), 'submission_revision':self.view['revision']}
        result = self.client.post('/api/jobs', json=body).json()
        jid, output = result['id'], Path(result['output_root'])
        folder = output/'01.match'/result['dataset']; folder.mkdir(parents=True)
        (folder/'sample_manifest.csv').write_text('sample_id,status\nA,OK\n')
        (output/'result.txt').write_text('test result')
        confirm = {'confirm':jid, 'output_root':str(output)}
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/delete',json=confirm).status_code,409)
        web.update_job(jid,status='STOPPED')
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/delete',json={}).status_code,400)
        result = self.client.post(f'/api/jobs/{jid}/delete',json=confirm)
        self.assertEqual(result.status_code,200,result.text)
        self.assertFalse(output.exists());self.assertTrue(self.raw.exists());self.assertTrue(self.source.exists())
        self.assertEqual(self.client.get(f'/api/jobs/{jid}').status_code,404)
        with web.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM actions WHERE job_id=?',(jid,)).fetchone()[0],0)

    def test_delete_rejects_shared_protected_and_failed_cleanup(self):
        jid=self.create();web.update_job(jid,status='STOPPED')
        def delete():
            return self.client.post(f'/api/jobs/{jid}/delete',json={'confirm':jid,'output_root':web.get_job_row(jid)['output_root']})
        self.assertEqual(delete().status_code,409)  # common default output root
        output=self.out/'exclusive';folder=output/'01.match'/'batch';folder.mkdir(parents=True)
        web.update_job(jid,output_root=str(output))
        sibling=self.create('other');web.update_job(sibling,status='STOPPED',output_root=str(output/'nested'))
        self.assertEqual(delete().status_code,409)
        web.update_job(sibling,output_root=str(self.out/'other'))
        with patch.object(web.shutil,'rmtree',side_effect=PermissionError('test denial')):
            self.assertEqual(delete().status_code,500)
        self.assertEqual(self.client.get(f'/api/jobs/{jid}').status_code,200)
        web.update_job(jid,output_root=str(self.raw))
        self.assertEqual(delete().status_code,409);self.assertTrue(self.raw.exists())
        unknown=self.root/'unowned';unknown.mkdir();(unknown/'keep').write_text('do not delete')
        web.update_job(jid,output_root=str(unknown))
        self.assertEqual(delete().status_code,409);self.assertTrue((unknown/'keep').exists())

    def test_stage_reports_pagination_download_and_isolation(self):
        jid=self.create(pipeline='ir_split')
        path=self.out/'03.IR_split_output'/'batch'/'ir_split_summary.csv';path.parent.mkdir(parents=True)
        path.write_text('sample_id,total_reads,matched_reads,matched_pct\nA,100,80,80.00\nB,100,0,0.00\n',encoding='utf-8')
        panda=self.out/'05.pandaseq'/'batch'/'pandaseq_summary.csv';panda.parent.mkdir(parents=True)
        panda.write_text('Sample,Total_Reads,OK_Reads,Merged_Percent\nA,80,70,87.50\n')
        reports=self.client.get(f'/api/jobs/{jid}/artifacts').json()['reports']
        self.assertTrue(next(r for r in reports if r['kind']=='split')['exists'])
        data=self.client.get(f'/api/jobs/{jid}/stage-summary',params={'kind':'split','limit':1,'offset':1}).json()
        self.assertEqual(data['total'],2);self.assertEqual(data['rows'][0]['sample_id'],'B')
        download=self.client.get(f'/api/jobs/{jid}/download?kind=pandaseq')
        self.assertEqual(download.status_code,200);self.assertIn('87.50',download.text)
        self.assertEqual(self.client.get(f'/api/jobs/{jid}/download?kind=../../raw').status_code,404)
        ten=self.create('ten',pipeline='10x_split')
        pre=self.out/'03.prefilter_data'/'ten'/'r1_r2_prefilter_summary.csv';pre.parent.mkdir(parents=True)
        pre.write_text('sample,total_pairs\nT,123\n')
        data=self.client.get(f'/api/jobs/{ten}/stage-summary?kind=prefilter').json()
        self.assertEqual(data['rows'][0]['total_pairs'],'123')
        self.assertEqual(self.client.get(f'/api/jobs/{ten}/stage-summary?kind=split').json()['total'],0)

    def test_project_relative_barcode_defaults(self):
        project=self.root/'project';reference=project/'reference';reference.mkdir(parents=True)
        barcode=reference/'8bp_barcodes.csv';barcode.write_text('name,sequence\n1,ACGTACGT\n')
        with patch.object(web,'PIPELINE_ROOT',project):
            defaults=self.client.get('/api/defaults').json()
            self.assertEqual(defaults['pipeline_root'],str(project))
            self.assertEqual(defaults['barcode_csv'],str(barcode))
            self.assertEqual(web.validate_barcode(None,True),barcode)
            self.assertIsNone(web.validate_barcode(None,False))
            listing=self.client.get('/api/browse',params={'kind':'barcode','path':str(reference)})
            self.assertEqual(listing.status_code,200,listing.text)
            self.assertEqual(listing.json()['entries'][0]['name'],'8bp_barcodes.csv')
            other=self.root/'outside.csv';other.write_text('not a reference')
            with self.assertRaises(web.HTTPException):web.validate_barcode(str(other),True)
            barcode.unlink()
            with self.assertRaises(web.HTTPException):web.validate_barcode(None,True)


if __name__ == '__main__': unittest.main()
