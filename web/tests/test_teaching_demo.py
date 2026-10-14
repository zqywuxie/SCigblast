"""Check the published teaching kit against actual input consumers."""
import csv
import gzip
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
KIT = ROOT / "web/static/demos/scigblast-teaching-v1.zip"


class TeachingDemoTests(unittest.TestCase):
    def test_published_inputs_convert_and_match(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            with zipfile.ZipFile(KIT) as archive:
                self.assertEqual(len(archive.namelist()), 13)
                archive.extractall(root)
            kit = root / "scigblast-teaching-v1"
            for path in (kit / "fastq/input").glob("*.gz"):
                lines = gzip.decompress(path.read_bytes()).decode().splitlines()
                self.assertEqual(len(lines), 80)
                for i in range(0, len(lines), 4):
                    self.assertEqual(len(lines[i + 1]), 150)
                    self.assertEqual(len(lines[i + 3]), 150)
            converted = root / "converted"
            result = subprocess.run([sys.executable, str(ROOT / "mapping/fasta.py"), "--input-dir", str(kit / "remapping/input"), "--output-dir", str(converted), "--chains", "IGH", "IGK"], capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONUTF8": "1"})
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(len(list(converted.glob("*.fasta"))), 3)
            self.assertEqual(sum(p.read_text().count(">") for p in converted.glob("*.fasta")), 6)
            for key, folder in [("ir_split", "IR_split"), ("10x_split", "10X_split"), ("igblast_base", "Igblast_base"), ("pig_igblast", "PigIgblast")]:
                with self.subTest(pipeline=key):
                    workbook = kit / "submission" / f"{key}.xlsx"
                    wb = load_workbook(workbook)
                    for row in wb.active.iter_rows(min_row=2):
                        row[5].value = str(kit / "fastq/input")
                    wb.save(workbook)
                    wb.close()
                    summary = root / key / "summary.csv"
                    env = {**os.environ, "PYTHONUTF8": "1", "SCIGBLAST_MATCH_INPUT": str(kit / "fastq/input"), "SCIGBLAST_RAW_INPUT_DIR": str(kit / "fastq/input"), "RAW_INPUT_DIR": str(kit / "fastq/input"), "SCIGBLAST_SUBMISSION_PATHS": str(workbook), "SCIGBLAST_SUBMISSION_XLSX": str(workbook), "SUBMISSION_XLSX": str(workbook), "SCIGBLAST_BARCODE_CSV": str(kit / "barcode.csv"), "SCIGBLAST_MATCH_OUTPUT": str(summary), "SCIGBLAST_OUTPUT_ROOT": str(root / key), "OUTPUT_ROOT": str(root / key)}
                    result = subprocess.run([sys.executable, str(ROOT / folder / "pipeline/01.match_sample.py")], capture_output=True, text=True, encoding="utf-8", env=env)
                    self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                    with summary.open(encoding="utf-8-sig") as handle:
                        rows = list(csv.DictReader(handle))
                    self.assertEqual(len(rows), 2)
                    self.assertEqual({row["sample_id"] for row in rows}, {"Demo_A", "Demo_B"})
                    self.assertTrue(all(row["status"] == "OK" for row in rows), rows)


if __name__ == "__main__":
    unittest.main()
