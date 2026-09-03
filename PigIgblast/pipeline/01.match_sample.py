#!/usr/bin/env python3
"""Match Pig FASTQ pairs to submission rows and emit the canonical manifest."""
from __future__ import annotations

import csv
import os
import posixpath
import re
import shlex
import sys
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - the XML fallback is used on lean servers
    load_workbook = None

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from tools.pipeline_config import load_config, path  # noqa: E402

CFG = load_config()
RAW = path(CFG, "RAW_INPUT_DIR", "")
SUBMISSION = path(CFG, "SUBMISSION_XLSX", "")
SUBMISSION_PATHS_VALUE = os.environ.get("SCIGBLAST_SUBMISSION_PATHS", "")
OUTPUT_ROOT = path(CFG, "SCIGBLAST_OUTPUT_ROOT", path(CFG, "OUTPUT_ROOT", str(HERE.parent / "output")).as_posix())
DATASET_LABEL = os.environ.get("SCIGBLAST_DATASET_LABEL", "").strip()
OUT = OUTPUT_ROOT / "01.match" / DATASET_LABEL if DATASET_LABEL else OUTPUT_ROOT / "01.match"
MATCH_OUTPUT = os.environ.get("SCIGBLAST_MATCH_OUTPUT", "").strip()
CONFIG_PATH = Path(os.environ.get("PIG_PIPELINE_CONFIG", HERE / "00.pipeline_config.env"))
FIELDS = [
    "note", "pair_id", "sample_id", "species", "chain_raw", "igblast_chains",
    "barcode_candidate", "barcode_name", "barcode_sequence", "dual_index",
    "match_method", "status", "error", "r1_path", "r2_path",
]
CHAIN_ORDER = ("TRA", "TRB", "TRD", "TRG", "IGH", "IGK", "IGL")
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _xml_text(element: ET.Element) -> str:
    return "".join(node.text or "" for node in element.iter(f"{{{NS_MAIN}}}t"))


def _column_number(cell_ref: str) -> int:
    match = re.match(r"[A-Za-z]+", cell_ref)
    if not match:
        return 0
    number = 0
    for char in match.group(0).upper():
        number = number * 26 + ord(char) - ord("A") + 1
    return number - 1


def _xml_sheet_paths(archive: zipfile.ZipFile) -> list[str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_map = {rel.attrib["Id"]: rel.attrib["Target"]
               for rel in rels.findall(f"{{{NS_REL}}}Relationship")}
    result: list[str] = []
    sheets = workbook.find(f"{{{NS_MAIN}}}sheets")
    if sheets is None:
        return result
    rel_id = f"{{{NS_DOC_REL}}}id"
    for sheet in sheets:
        target = rel_map.get(sheet.attrib.get(rel_id, ""))
        if target:
            result.append(posixpath.normpath(posixpath.join("xl", target)).lstrip("/"))
    return result


def _xml_sheet_rows(archive: zipfile.ZipFile, sheet_path: str, shared: list[str]) -> list[list[str]]:
    root = ET.fromstring(archive.read(sheet_path))
    sheet_data = root.find(f"{{{NS_MAIN}}}sheetData")
    if sheet_data is None:
        return []
    rows: list[list[str]] = []
    for row in sheet_data:
        values: dict[int, str] = {}
        max_col = -1
        for cell in row.findall(f"{{{NS_MAIN}}}c"):
            col = _column_number(cell.attrib.get("r", "A1")); max_col = max(max_col, col)
            kind = cell.attrib.get("t")
            if kind == "inlineStr":
                value = _xml_text(cell)
            else:
                node = cell.find(f"{{{NS_MAIN}}}v")
                value = node.text if node is not None else ""
                if kind == "s" and value:
                    value = shared[int(value)]
            values[col] = value
        rows.append([values.get(i, "") for i in range(max_col + 1)])
    return rows


def iter_sheet_rows(workbook: Path):
    """Yield worksheet rows using openpyxl when present, otherwise stdlib XML."""
    if load_workbook is not None:
        wb = load_workbook(workbook, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = ws.iter_rows(values_only=True)
                yield next(rows, None), rows
        finally:
            wb.close()
        return
    with zipfile.ZipFile(workbook) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [_xml_text(si) for si in root.findall(f"{{{NS_MAIN}}}si")]
        for sheet_path in _xml_sheet_paths(archive):
            rows = iter(_xml_sheet_rows(archive, sheet_path, shared))
            yield next(rows, None), rows


def config_array(name: str, default: tuple[str, ...]) -> tuple[str, ...]:
    env = os.environ.get(name + "_TEXT", "")
    if env:
        return tuple(item for item in env.splitlines() if item)
    try:
        text = CONFIG_PATH.read_text(encoding="utf-8-sig")
    except OSError:
        return default
    match = re.search(rf"(?ms)^\s*{re.escape(name)}\s*=\s*\((.*?)\)", text)
    if not match:
        return default
    values = tuple(shlex.split(match.group(1), comments=True, posix=True))
    return values or default


def read_rules() -> tuple[tuple[tuple[str, str], ...], tuple[str, ...]]:
    pair_values = config_array("SCIGBLAST_RAW_READ_SUFFIX_PAIRS", ("_R1|_R2", "_1|_2", "_j|_v"))
    ext_values = config_array("SCIGBLAST_RAW_FASTQ_EXTENSIONS", (".fq.gz", ".fastq.gz", ".fq", ".fastq"))
    pairs: list[tuple[str, str]] = []; seen: set[str] = set()
    for value in pair_values:
        if value.count("|") != 1:
            raise ValueError(f"invalid suffix pair {value!r}; expected R1|R2")
        r1, r2 = (item.strip() for item in value.split("|", 1))
        if not r1 or not r2 or r1.casefold() == r2.casefold():
            raise ValueError(f"invalid suffix pair {value!r}")
        for suffix in (r1, r2):
            if "/" in suffix or "\\" in suffix or suffix.casefold().endswith((".fq", ".fastq", ".fq.gz", ".fastq.gz")):
                raise ValueError(f"suffix must not contain path/FASTQ extension: {suffix!r}")
            if suffix.casefold() in seen:
                raise ValueError(f"duplicate read suffix: {suffix!r}")
            seen.add(suffix.casefold())
        pairs.append((r1, r2))
    extensions: list[str] = []
    for value in ext_values:
        if not value.startswith(".") or "/" in value or "\\" in value:
            raise ValueError(f"invalid FASTQ extension: {value!r}")
        if value.casefold() not in {item.casefold() for item in extensions}:
            extensions.append(value)
    return tuple(pairs), tuple(sorted(extensions, key=len, reverse=True))


SUFFIX_PAIRS, FASTQ_EXTENSIONS = read_rules()


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text).replace("\u3000", " ")).strip()


def key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", norm(value).casefold())


def normalize_chains(raw: str) -> tuple[str, ...]:
    compact = re.sub(r"[\s_-]+", "", unicodedata.normalize("NFKC", raw or "")).upper()
    compact = compact.replace("Α", "ALPHA").replace("Β", "BETA")
    aliases = {
        "7C": CHAIN_ORDER, "BOTH": CHAIN_ORDER,
        "T": CHAIN_ORDER[:4], "TCR": CHAIN_ORDER[:4],
        "B": CHAIN_ORDER[4:], "BCR": CHAIN_ORDER[4:],
        "AB": ("TRA", "TRB"), "ALPHABETA": ("TRA", "TRB"),
        "TCRALPHA": ("TRA",), "TCRA": ("TRA",),
        "TCRBETA": ("TRB",), "TCRB": ("TRB",),
    }
    if compact in aliases:
        return tuple(aliases[compact])
    found: set[str] = set()
    for token in filter(None, re.split(r"[,;/+|]+", compact)):
        if token in CHAIN_ORDER:
            found.add(token)
        elif token == "IGH*" or re.fullmatch(r"IGH[ADGME]", token):
            found.add("IGH")
        elif token in aliases:
            found.update(aliases[token])
    return tuple(chain for chain in CHAIN_ORDER if chain in found)


def workbook_paths(root: Path) -> list[Path]:
    if root.is_file() and root.suffix.casefold() == ".xlsx":
        return [root]
    if root.is_dir():
        files = sorted(p for p in root.rglob("*.xlsx") if p.is_file())
        if files:
            return files
    raise FileNotFoundError(f"No .xlsx submission workbook found: {root}")


def find_column(keys: dict[str, int], *names: str) -> int | None:
    for name in names:
        if key(name) in keys:
            return keys[key(name)]
    return None


def load_submission() -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    configured = [Path(item) for item in SUBMISSION_PATHS_VALUE.splitlines() if item.strip()] or [SUBMISSION]
    workbooks: list[Path] = []
    for root in configured:
        workbooks.extend(workbook_paths(root))
    for workbook in list(dict.fromkeys(workbooks)):
        for header, rows in iter_sheet_rows(workbook):
            if not header:
                continue
            keys = {key(value): i for i, value in enumerate(header) if value is not None}
            sample_i = find_column(keys, "Sample ID File new name", "Sample ID")
            dual_i = find_column(keys, "Dual Index", "Dual Index（测序条码）候选")
            species_i = find_column(keys, "Species", "Speicies")
            chain_i = find_column(keys, "Chain")
            note_i = find_column(keys, "Note", "Note（数据文件储存地址）")
            if None in (sample_i, dual_i, species_i, chain_i):
                continue
            current_note = ""
            for values in rows:
                if not values or not any(value not in (None, "") for value in values):
                    continue
                if note_i is not None and note_i < len(values) and norm(values[note_i]):
                    current_note = norm(values[note_i])
                sample = norm(values[sample_i])
                if not sample or key(sample) in {"sampleid", "sampleidfilenewname"}:
                    continue
                records.append({"note": current_note, "sample_id": sample,
                                "dual_index": norm(values[dual_i]).upper(), "species": norm(values[species_i]),
                                "chain_raw": norm(values[chain_i])})
    if not records:
        raise ValueError("submission contains no usable Pig sample rows")
    return records


def note_matches(file_path: Path, note: str) -> bool:
    if not note:
        return False
    file_norm = os.path.normcase(os.path.normpath(str(file_path.resolve())))
    for value in re.split(r"[;\n]+", note):
        note_norm = os.path.normcase(os.path.normpath(value.strip()))
        if note_norm and (file_norm == note_norm or file_norm.startswith(note_norm.rstrip(os.sep) + os.sep)):
            return True
    return False


def parse_dual(name: str) -> str:
    match = re.search(r"(?:^|[_-])([A-Ha-h]\d{2})(?=_L\d{2}(?:[_-]|\.|$))", name)
    return match.group(1).upper() if match else ""


def read_matches(name: str) -> list[tuple[str, str]]:
    lower = name.casefold(); out: list[tuple[str, str]] = []
    for ext in FASTQ_EXTENSIONS:
        if not lower.endswith(ext.casefold()):
            continue
        core = name[:-len(ext)]; folded = core.casefold()
        for r1_suffix, r2_suffix in SUFFIX_PAIRS:
            for side, suffix in (("R1", r1_suffix), ("R2", r2_suffix)):
                if folded.endswith(suffix.casefold()):
                    out.append((side, core[:-len(suffix)])); continue
                lane = re.search(re.escape(suffix.casefold()) + r"([._-]\d+)$", folded)
                if lane:
                    out.append((side, core[:lane.start()] + lane.group(1)))
        break
    return out


def pair_fastqs() -> tuple[list[dict[str, object]], int]:
    grouped: dict[tuple[Path, str], dict[str, list[Path]]] = defaultdict(lambda: {"R1": [], "R2": []})
    ambiguous: set[tuple[Path, str]] = set(); ignored = 0
    for file_path in sorted(p for p in RAW.rglob("*") if p.is_file()):
        matches = read_matches(file_path.name)
        if not matches:
            ignored += 1; continue
        keys = {(file_path.parent, pair) for _, pair in matches}
        if len(matches) != 1 or len(keys) != 1:
            ambiguous.update(keys); continue
        side, pair = matches[0]; grouped[(file_path.parent, pair)][side].append(file_path.resolve())
    jobs: list[dict[str, object]] = []
    for (parent, pair), sides in sorted(grouped.items(), key=lambda item: (str(item[0][0]), item[0][1])):
        error = "filename matches multiple configured read rules" if (parent, pair) in ambiguous else ""
        if len(sides["R1"]) != 1 or len(sides["R2"]) != 1:
            error = error or ("missing R1 pair" if not sides["R1"] else "missing R2 pair" if not sides["R2"] else "multiple FASTQs resolve to one pair")
        jobs.append({"pair": pair, "r1": sides["R1"][0] if len(sides["R1"]) == 1 else None,
                     "r2": sides["R2"][0] if len(sides["R2"]) == 1 else None, "error": error})
    return jobs, ignored


def empty_row(pair: str, r1: Path | None, r2: Path | None) -> dict[str, str]:
    return {field: "" for field in FIELDS} | {"pair_id": pair, "status": "ERROR",
        "r1_path": str(r1) if r1 else "", "r2_path": str(r2) if r2 else ""}


def main() -> int:
    configured_submissions = [Path(item) for item in SUBMISSION_PATHS_VALUE.splitlines() if item.strip()] or [SUBMISSION]
    if not RAW.is_dir() or not any(item.exists() for item in configured_submissions):
        print(f"[PIG][match] missing input/submission: {RAW} {SUBMISSION}", file=sys.stderr); return 2
    try:
        submission = load_submission(); pairs, ignored = pair_fastqs()
    except Exception as exc:
        print(f"[PIG][match] {exc}", file=sys.stderr); return 2
    rows: list[dict[str, str]] = []
    for job in pairs:
        pair, r1, r2 = str(job["pair"]), job["r1"], job["r2"]
        row = empty_row(pair, r1 if isinstance(r1, Path) else None, r2 if isinstance(r2, Path) else None)
        if job["error"]:
            row["error"] = str(job["error"]); rows.append(row); continue
        assert isinstance(r1, Path) and isinstance(r2, Path)
        dual = parse_dual(r1.name); row["dual_index"] = dual
        if not dual:
            row["error"] = "Dual Index not found in filename"; rows.append(row); continue
        has_notes = any(record["note"] for record in submission)
        batch = [record for record in submission if note_matches(r1, record["note"])] if has_notes else submission
        if not batch:
            row["error"] = "no submission Note matches FASTQ path"; rows.append(row); continue
        selected = [record for record in batch if record["dual_index"] == dual]
        if not selected:
            row["error"] = f"no Dual Index match within the Note batch: {dual}"; rows.append(row); continue
        if any(record["species"].casefold() != "pig" for record in selected):
            row["error"] = "matching submission row is not Species=Pig"; rows.append(row); continue
        by_sample: dict[str, list[dict[str, str]]] = defaultdict(list)
        for record in selected:
            by_sample[record["sample_id"]].append(record)
        if len(by_sample) != 1:
            row["error"] = "unsplittable physical pair maps to multiple samples: " + ", ".join(sorted(by_sample)); rows.append(row); continue
        sample, records = next(iter(by_sample.items()))
        chain_raw = ",".join(dict.fromkeys(record["chain_raw"] for record in records if record["chain_raw"]))
        chains = tuple(chain for chain in CHAIN_ORDER if any(chain in normalize_chains(record["chain_raw"]) for record in records))
        if not chains:
            row["error"] = "Chain is empty or unsupported"; rows.append(row); continue
        row.update({"note": records[0]["note"], "sample_id": sample, "species": records[0]["species"],
                    "chain_raw": chain_raw, "igblast_chains": ",".join(chains),
                    "match_method": "dual_index+note" if has_notes else "dual_index", "status": "OK", "error": ""})
        rows.append(row)
    OUT.mkdir(parents=True, exist_ok=True)
    destinations = [OUT / "sample_manifest.csv", OUT / "match_summary.csv"]
    if MATCH_OUTPUT:
        destinations.insert(0, Path(MATCH_OUTPUT))
    for destination in destinations:
        destination.parent.mkdir(parents=True, exist_ok=True)
        temp = destination.with_suffix(destination.suffix + ".tmp")
        with temp.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS); writer.writeheader(); writer.writerows(rows)
        os.replace(temp, destination)
    ok = sum(row["status"] == "OK" for row in rows)
    print(f"[PIG][match] completed={len(rows)} failed={len(rows)-ok} total={len(rows)} usable={ok} ignored_files={ignored}")
    return 0 if ok and (not len(rows)-ok or CFG.get("ALLOW_PARTIAL") == "1") else 1


if __name__ == "__main__":
    raise SystemExit(main())
